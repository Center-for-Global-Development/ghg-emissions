"""
05_extract_gcp.py
Extract Global Carbon Project (GCB 2025) data for CO2 measures.

Inputs:
  National_Fossil_Carbon_Emissions_2025_v0.3.xlsx  (sheet: Territorial Emissions)
  National_LandUseChange_Carbon_Emissions_2025v0.2.xlsx  (sheets: BLUE, OSCAR, LUCE)

Output: data/outputs/gcp_long.csv

Standard schema: source, iso_code, year, measure, value_Mt

Sources produced:
  GCP_fossil  -> co2_excLUC  (fossil CO2 only)
  GCP_BLUE    -> co2_incLUC  (fossil + BLUE LUC model)
  GCP_OSCAR   -> co2_incLUC  (fossil + OSCAR LUC model)
  GCP_LUCE    -> co2_incLUC  (fossil + LUCE LUC model)

Note: GCP provides CO2 only (no all-GHG total).

Unit conversion: raw values in MtC/yr.  value_Mt (MtCO2) = raw_value * 3.664

Country name -> ISO3 mapping:
  GCP files use country names (not ISO codes) as column headers.
  Mapping is built from data/country_groups.csv (iso_code, country_name).
  Country names are matched case-insensitively.
  Unmatched names are logged and dropped.

Coverage: 1850-2024 (all four series — GCB 2025 v0.3 / v0.2)

Exclusions:
  - "World" / "WORLD" totals excluded
  - Bunker fuels column (present in fossil file) excluded
  - Statistical difference column excluded
"""

import os
import re
import pandas as pd
import openpyxl

SCRIPTS_DIR = os.path.dirname(os.path.abspath(__file__))
QA_DIR      = os.path.dirname(SCRIPTS_DIR)
RAW_DIR     = os.path.join(QA_DIR, "data", "raw")

FOSSIL_FILE = os.path.join(RAW_DIR, "National_Fossil_Carbon_Emissions_2025_v0.3.xlsx")
LUC_FILE    = os.path.join(RAW_DIR, "National_LandUseChange_Carbon_Emissions_2025v0.2.xlsx")
CG_FILE     = os.path.join(QA_DIR, "data", "country_groups.csv")
OUT         = os.path.join(QA_DIR, "data", "outputs", "gcp_long.csv")

GCP_CONV = 3.664   # MtC -> MtCO2

# ---------------------------------------------------------------------------
# Name -> ISO3 mapping from country_groups.csv
# ---------------------------------------------------------------------------
cg = pd.read_csv(CG_FILE)[["iso_code", "country_name"]]
name_to_iso = {
    row.country_name.strip().upper(): row.iso_code
    for _, row in cg.iterrows()
}

# Manual overrides for names that differ between GCP and the Combined workbook
GCP_NAME_OVERRIDES = {
    "CHINA, P.R. (MAINLAND)":    "CHN",
    "CHINA, MAINLAND":           "CHN",
    "CHINESE TAIPEI":            "TWN",
    "IRAN":                      "IRN",
    "SOUTH KOREA":               "KOR",
    "NORTH KOREA":               "PRK",
    "RUSSIA":                    "RUS",
    "BOLIVIA":                   "BOL",
    "VIETNAM":                   "VNM",
    "TANZANIA":                  "TZA",
    "MOLDOVA":                   "MDA",
    "CZECH REPUBLIC":            "CZE",
    "DEMOCRATIC REPUBLIC OF CONGO": "COD",
    "DR CONGO":                  "COD",
    "IVORY COAST":               "CIV",
    "BRUNEI":                    "BRN",
    "NORTH MACEDONIA":           "MKD",
    "SOUTH SUDAN":               "SSD",
    "ESWATINI":                  "SWZ",
    "TIMOR":                     "TLS",
    "EAST TIMOR":                "TLS",
    "VENEZUELA":                 "VEN",
    "LAOS":                      "LAO",
    "MYANMAR":                   "MMR",
    "BURMA":                     "MMR",
    "SYRIA":                     "SYR",
    "CAPE VERDE":                "CPV",
    "FEDERATED STATES OF MICRONESIA": "FSM",
    "MICRONESIA":                "FSM",
    "SAINT KITTS AND NEVIS":     "KNA",
    "SAINT LUCIA":               "LCA",
    "SAINT VINCENT AND THE GRENADINES": "VCT",
    "TRINIDAD AND TOBAGO":       "TTO",
    "ANTIGUA AND BARBUDA":       "ATG",
    "SAO TOME AND PRINCIPE":     "STP",
    "COMOROS":                   "COM",
    "BAHAMAS":                   "BHS",
    "BARBADOS":                  "BRB",
    "BERMUDA":                   "BMU",
    "COOK ISLANDS":              "COK",
    "FAROE ISLANDS":             "FRO",
    "FALKLAND ISLANDS":          "FLK",
    "FRENCH POLYNESIA":          "PYF",
    "FRENCH GUIANA":             "GUF",
    "GREENLAND":                 "GRL",
    "GUADELOUPE":                "GLP",
    "MARTINIQUE":                "MTQ",
    "MAYOTTE":                   "MYT",
    "NEW CALEDONIA":             "NCL",
    "NIUE":                      "NIU",
    "PALAU":                     "PLW",
    "REUNION":                   "REU",
    "RYUKYU ISLANDS":            "JPN",   # now part of Japan
    "SAINT PIERRE AND MIQUELON": "SPM",
    "WALLIS AND FUTUNA ISLANDS": "WLF",
    "WESTERN SAHARA":            "ESH",
    # Additional overrides for GCP fossil file naming conventions
    "ANTIGUA & BARBUDA":         "ATG",
    "BONAIRE, SAINT EUSTATIUS, AND SABA": "BES",
    "BOSNIA & HERZEGOVINA":      "BIH",
    "BRUNEI (DARUSSALAM)":       "BRN",
    "CHINA (MAINLAND)":          "CHN",
    "CONGO":                     "COG",
    "COTE D IVOIRE":             "CIV",
    "CURACAO":                   "CUW",
    "DEMOCRATIC PEOPLE S REPUBLIC OF KOREA": "PRK",
    "DEMOCRATIC REPUBLIC OF THE CONGO (FORMERLY ZAIRE)": "COD",
    "DEMOCRATIC REPUBLIC OF THE CONGO": "COD",
    "EGYPT":                     "EGY",
    "FAEROE ISLANDS":            "FRO",
    "FRANCE (INCLUDING MONACO)": "FRA",
    "GAMBIA":                    "GMB",
    "GUINEA BISSAU":             "GNB",
    "GUINEA-BISSAU":             "GNB",
    "HONG KONG SPECIAL ADMINSTRATIVE REGION OF CHINA": "HKG",
    "HONG KONG":                 "HKG",
    "ISLAMIC REPUBLIC OF IRAN":  "IRN",
    "ITALY (INCLUDING SAN MARINO)": "ITA",
    "KYRGYZSTAN":                "KGZ",
    "LAO PEOPLE S DEMOCRATIC REPUBLIC": "LAO",
    "LAO PDR":                   "LAO",
    "LIBYAN ARAB JAMAHIRIYAH":   "LBY",
    "LIBYA":                     "LBY",
    "MACAU SPECIAL ADMINSTRATIVE REGION OF CHINA": "MAC",
    "MACAO":                     "MAC",
    "MACEDONIA":                 "MKD",
    "MICRONESIA (FEDERATED STATES OF)": "FSM",
    "MYANMAR (FORMERLY BURMA)":  "MMR",
    "OCCUPIED PALESTINIAN TERRITORY": "PSE",
    "STATE OF PALESTINE":        "PSE",
    "PLURINATIONAL STATE OF BOLIVIA": "BOL",
    "REPUBLIC OF CAMEROON":      "CMR",
    "REPUBLIC OF KOREA":         "KOR",
    "REPUBLIC OF MOLDOVA":       "MDA",
    "REPUBLIC OF SOUTH SUDAN":   "SSD",
    "REPUBLIC OF SUDAN":         "SDN",
    "SUDAN":                     "SDN",
    "RUSSIAN FEDERATION":        "RUS",
    "SAINT MARTIN (DUTCH PORTION)": "SXM",
    "SAO TOME & PRINCIPE":       "STP",
    "ST. KITTS-NEVIS":           "KNA",
    "ST. PIERRE & MIQUELON":     "SPM",
    "ST. VINCENT & THE GRENADINES": "VCT",
    "TAIWAN":                    "TWN",
    "TIMOR-LESTE (FORMERLY EAST TIMOR)": "TLS",
    "TIMOR-LESTE":               "TLS",
    "T\u00dcRKIYE":              "TUR",   # Türkiye (uppercase)
    "T\u00dcRKEY":               "TUR",   # Turkey (uppercase, fallback)
    "TURKEY":                    "TUR",
    "UNITED REPUBLIC OF TANZANIA": "TZA",
    "UNITED STATES OF AMERICA":  "USA",
    "USA":                       "USA",
    "VIET NAM":                  "VNM",
    "YEMEN":                     "YEM",
    # LUC file uses Côte d'Ivoire with Unicode
    "C\u00d4TE D'IVOIRE":        "CIV",
    "COTE D'IVOIRE":             "CIV",
}
name_to_iso.update(GCP_NAME_OVERRIDES)

NON_COUNTRY_NAMES = {
    "WORLD", "BUNKER FUELS", "STATISTICAL DIFFERENCE", "INTERNATIONAL",
    "OCEAN SINK", "LAND SINK", "GLOBAL OCEAN", "GLOBAL LAND",
    # GCP fossil aggregate columns
    "KP ANNEX B", "NON KP ANNEX B", "OECD", "NON-OECD", "EU27",
    "AFRICA", "ASIA", "CENTRAL AMERICA", "EUROPE", "MIDDLE EAST",
    "NORTH AMERICA", "OCEANIA", "SOUTH AMERICA",
    "INTERNATIONAL SHIPPING", "INTERNATIONAL AVIATION",
}


def gcp_col_to_iso(col_name):
    """Map a GCP column header (country name) to ISO3, or None if not found."""
    upper = str(col_name).strip().upper()
    if upper in NON_COUNTRY_NAMES or not upper:
        return None
    return name_to_iso.get(upper)


def read_gcp_wide(path, sheet_name, data_start_row, name_row):
    """
    Read a GCP-format sheet (rows=years, columns=countries).
    data_start_row and name_row are 1-based row numbers.
    Returns a long DataFrame: iso_code, year, raw_value (MtC).
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name]

    # Read country name row
    col_names = {}
    for row in ws.iter_rows(min_row=name_row, max_row=name_row):
        for cell in row:
            if not hasattr(cell, "column") or cell.column is None or cell.column == 1:
                continue   # skip EmptyCell objects and col A (year)
            if cell.value is not None:
                col_names[cell.column] = cell.value

    # Build iso mapping for this sheet
    iso_map = {}
    unmatched = []
    for col_idx, name in col_names.items():
        iso = gcp_col_to_iso(name)
        if iso:
            iso_map[col_idx] = iso
        else:
            upper = str(name).strip().upper()
            if upper not in NON_COUNTRY_NAMES:
                unmatched.append(name)

    if unmatched:
        print(f"    UNMATCHED names in {os.path.basename(path)} [{sheet_name}]: "
              f"{unmatched[:10]}{'...' if len(unmatched)>10 else ''}")

    # Read data rows
    rows = []
    for r in ws.iter_rows(min_row=data_start_row, values_only=True):
        year = r[0]
        if not isinstance(year, (int, float)):
            continue
        year = int(year)
        for col_idx, iso in iso_map.items():
            val = r[col_idx - 1]   # 0-based index
            if val is not None:
                rows.append({"iso_code": iso, "year": year, "raw_value": val})

    wb.close()
    return pd.DataFrame(rows)


# ---------------------------------------------------------------------------
# 1. Fossil CO2
# ---------------------------------------------------------------------------
print("Reading GCP fossil ...")
# Fossil: row 11 = uppercase country names, row 12 = proper names, data from row 13
fossil_df = read_gcp_wide(FOSSIL_FILE, "Territorial Emissions",
                          data_start_row=13, name_row=11)
fossil_df["source"]   = "GCP_fossil"
fossil_df["measure"]  = "co2_excLUC"
fossil_df["value_Mt"] = fossil_df["raw_value"] * GCP_CONV
print(f"  Fossil: {len(fossil_df):,} rows, {fossil_df['iso_code'].nunique()} countries")

# ---------------------------------------------------------------------------
# 2. LUC CO2 — three model sheets
# ---------------------------------------------------------------------------
LUC_SHEETS = {
    "BLUE":  "GCP_BLUE",
    "OSCAR": "GCP_OSCAR",
    "LUCE":  "GCP_LUCE",
}

luc_chunks = []
for sheet, source_label in LUC_SHEETS.items():
    print(f"Reading GCP LUC [{sheet}] ...")
    # LUC: row 8 = country names, data from row 9
    luc = read_gcp_wide(LUC_FILE, sheet, data_start_row=9, name_row=8)
    luc["source"] = source_label
    luc["measure"] = "co2_incLUC_luc_only"   # temporary; will add fossil below
    luc["value_Mt"] = luc["raw_value"] * GCP_CONV
    print(f"  {source_label}: {len(luc):,} rows, {luc['iso_code'].nunique()} countries")
    luc_chunks.append(luc)

# ---------------------------------------------------------------------------
# 3. Combine: GCP_fossil = co2_excLUC; GCP_fossil + LUC = co2_incLUC
# ---------------------------------------------------------------------------
fossil_for_merge = fossil_df[["iso_code", "year", "value_Mt"]].rename(
    columns={"value_Mt": "fossil_Mt"}
)

chunks = [fossil_df[["source", "iso_code", "year", "measure", "value_Mt"]]]

for luc in luc_chunks:
    source_label = luc["source"].iloc[0]
    luc_vals = luc[["iso_code", "year", "value_Mt"]].rename(columns={"value_Mt": "luc_Mt"})
    # Outer join so countries present in fossil but absent from LUC (e.g. TWN, HKG,
    # territories) are kept with luc_Mt=0 rather than silently dropped.
    merged = fossil_for_merge.merge(luc_vals, on=["iso_code", "year"], how="outer")
    merged["fossil_Mt"] = merged["fossil_Mt"].fillna(0)
    merged["luc_Mt"]    = merged["luc_Mt"].fillna(0)
    merged["value_Mt"]  = merged["fossil_Mt"] + merged["luc_Mt"]
    merged["source"]    = source_label
    merged["measure"]   = "co2_incLUC"
    chunks.append(merged[["source", "iso_code", "year", "measure", "value_Mt"]])

result = pd.concat(chunks, ignore_index=True)
result = result.dropna(subset=["value_Mt"])

os.makedirs(os.path.dirname(OUT), exist_ok=True)
result.to_csv(OUT, index=False)
print(
    f"Written: gcp_long.csv  "
    f"({len(result):,} rows, sources: {sorted(result['source'].unique())}, "
    f"measures: {sorted(result['measure'].unique())})"
)
